"""Import the published RVO workbook without geocoding or inferring capacity.

Requires openpyxl. Pass an existing workbook path or download the stated release.
"""
import collections
import datetime
import hashlib
import json
from pathlib import Path
import sys
import urllib.request
import openpyxl

ROOT=Path(__file__).resolve().parents[1]
URL='https://www.rvo.nl/sites/default/files/2026-09/SDE-projecten-in-beheer-juli-2026.xlsx'
FIELDS=['SDE ronde','Referentie','Categorie','Thema','Aanvrager Naam','Adres','Postcode','Plaats lokatie','Gemeente','RES-regio','Provincie','Looptijd [jr.]','Maximale subsidie [€]','Opstelling zon-pv','Gerealiseerd','Ingebruikname jaar']

def main():
    raw=Path(sys.argv[1]).read_bytes() if len(sys.argv)>1 else urllib.request.urlopen(URL,timeout=60).read()
    archive=ROOT/'research/raw/rvo-sde-july-2026.xlsx';archive.write_bytes(raw)
    sheet=openpyxl.load_workbook(archive,read_only=True,data_only=True).active
    header=[str(x).strip() for x in next(sheet.iter_rows(min_row=11,max_row=11,values_only=True))[1:17]]
    if header!=FIELDS: raise ValueError(f'RVO schema changed: {header}')
    projects=[];seen=set();all_count=0
    for row in sheet.iter_rows(min_row=12,values_only=True):
        if not row[2]: continue
        all_count+=1
        if row[3] not in ['Zon','Wind op land','Wind op zee']: continue
        ref=str(row[2]).strip()
        if ref in seen: raise ValueError(f'Duplicate RVO project reference: {ref}')
        seen.add(ref)
        def val(index):
            return str(row[index]).strip() if row[index] is not None else ''
        year=None if row[16] == '(Nog) niet' else row[16]
        if year is not None and (not isinstance(year,(int,float)) or not 1900<=year<=2100):
            raise ValueError(f'Unexpected commissioning year: {ref} {year}')
        projects.append({'id':ref,'round':val(1),'technology':{'Zon':'Solar','Wind op land':'Onshore wind','Wind op zee':'Offshore wind'}[val(3)],
                         'theme':val(4),'applicant':val(5),'address':val(6),'postcode':val(7),'place':val(8),'municipality':val(9),
                         'province':val(11),'installation':{'Daksysteem':'Rooftop','Veldsysteem':'Ground-mounted','Watersysteem':'Floating'}.get(val(14),'Unspecified'),
                         'realized':True if val(15)=='Ja' else False if val(15)=='Nee' else None,
                         'commissioningYear':int(year) if year is not None else None,'commissioningNote':val(16) if year is None else ''})
    projects.sort(key=lambda p:(p['municipality'],p['applicant'],p['id']))
    meta={'source':'RVO SDE(+)(+)','referenceDate':'2026-07','retrievedAt':datetime.datetime.now(datetime.timezone.utc).isoformat(),
          'sourceUrl':URL,'sha256':hashlib.sha256(raw).hexdigest(),'workbookRecords':all_count,'count':len(projects),
          'technologies':dict(collections.Counter(p['technology'] for p in projects)),
          'solarInstallationTypes':dict(collections.Counter(p['installation'] for p in projects if p['technology']=='Solar')),
          'scope':'Solar and onshore/offshore wind subsidy records only. Other workbook technologies excluded.',
          'limitations':'Not a census of physical assets. No capacity column, map coordinates or telemetry. Applicant is not necessarily current operator. Anonymisation is preserved; no identity reconstruction or automatic OSM joins.',
          'license':'RVO public project publication; original workbook notes and anonymisation retained.'}
    (ROOT/'public/data/rvo-projects.json').write_text(json.dumps({'meta':meta,'projects':projects},ensure_ascii=False,separators=(',',':')))
    for p in [ROOT/'data/rvo-projects-meta.json',ROOT/'public/data/rvo-projects-meta.json']:
        p.write_text(json.dumps(meta,indent=2,ensure_ascii=False)+'\n')
    print(json.dumps(meta,indent=2))

if __name__=='__main__':main()
